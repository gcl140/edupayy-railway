# from openpyxl import Workbook
# from .tables import InvoiceTable
# from django_tables2 import RequestConfig
# from django.core.mail import send_mail
# from django.urls import reverse
# from .forms import ParentContactForm
# from django.contrib.auth import get_user_model
# import csv
# from django.contrib import messages
# from email import message
# from django.contrib.admin.views.decorators import staff_member_required
# from datetime import datetime, date
# from django.utils.timezone import now
# from django.utils import timezone
# from decimal import Decimal, ROUND_HALF_UP
# import stripe
# from django.conf import settings
# from django.views.decorators.csrf import csrf_exempt
# from django.shortcuts import render, redirect, get_object_or_404
# from django.http import HttpResponse, JsonResponse
# from django.template.loader import render_to_string
# from weasyprint import HTML
# from .forms import FeeStructureForm, InvoiceForm, StudentProfileForm
# from .models import Payment, StudentProfile, ParentRequest, Invoice, PaymentMethod, FeeStructure, InvoiceFee
# from yuzzaz.models import Notification, UserNotificationStatus
# from django.contrib.auth.decorators import login_required
# from django.db.models import Q
import csv
import stripe
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.contrib.admin.views.decorators import staff_member_required
from django.core.mail import send_mail
from django.db.models import Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt

from django_tables2 import RequestConfig
from openpyxl import Workbook
from weasyprint import HTML
from django.utils.timezone import now

from .forms import ParentContactForm, FeeStructureForm, InvoiceForm, StudentProfileForm
from .models import (
    StudentProfile, ParentRequest, Payment, Invoice,
    PaymentMethod, FeeStructure, InvoiceFee
)
from .tables import InvoiceTable
from yuzzaz.models import Notification, UserNotificationStatus

# views.py
from django.shortcuts import render, redirect
from django.contrib import messages
from .forms import MultiInvoiceForm
from .models import Invoice, InvoiceFee
from functools import wraps
from django.shortcuts import render
from .models import StudentProfile, FeeStructure
from .forms import InvoiceForm

User = get_user_model()
def parents_required(view_func):
    @wraps(view_func)
    def _wrapped_view(request, *args, **kwargs):
        if request.user.is_authenticated and not request.user.is_parent:
            return render(request, "yuzzaz/403.html")
        return view_func(request, *args, **kwargs)
    return _wrapped_view


@login_required
def download_receipt(request, payment_id):
    payment = Payment.objects.select_related('invoice__student__user').get(id=payment_id)

    html_string = render_to_string("billing/receipt_template.html", {
        "payment": payment
    })

    html = HTML(string=html_string)
    pdf = html.write_pdf()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="receipt-{payment.id}.pdf"'
    return response


@login_required
def download_invoice_pdf(request, invoice_id):
    invoice = Invoice.objects.select_related('student__user').prefetch_related('invoice_fees__fee_structure').get(id=invoice_id)

    html_string = render_to_string("billing/invoice_template.html", {
        "invoice": invoice,
        "student": invoice.student,
        "fees": invoice.invoice_fees.all(),
    })

    pdf = HTML(string=html_string).write_pdf()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="invoice-{invoice.id}.pdf"'
    return response


@login_required
def payments_history(request):
    user = request.user
    query = request.GET.get('q')

    if user.is_staff:
        all_payments = Payment.objects.select_related('invoice', 'invoice__student').order_by('-payment_date')
        students = StudentProfile.objects.all()
        payment_methods = None

    elif user.is_parent:
        all_payments = Payment.objects.filter(
            invoice__student__parent=user
        ).select_related('invoice', 'invoice__student').order_by('-payment_date')
        students = StudentProfile.objects.filter(parent=user)
        payment_methods = user.payment_methods.all()

    else:
        profile = StudentProfile.objects.filter(user=user).first()
        if not profile:
            messages.warning(request, "Student profile not found. Please contact admin.")
            return redirect('student_dashboard')

        all_payments = Payment.objects.filter(
            invoice__student=profile
        ).select_related('invoice', 'invoice__student').order_by('-payment_date')
        students = [profile]
        payment_methods = None

    if query:
        all_payments = all_payments.filter(
            Q(invoice__student__user__first_name__icontains=query) |
            Q(invoice__student__user__last_name__icontains=query) |
            Q(method__icontains=query) |
            Q(transaction_id__icontains=query)
        ).distinct()

    context = {
        'user': user,
        'students': students,
        'methods': payment_methods,
        'recent_payments': all_payments,
        'query': query,
    }
    return render(request, 'billing/payments_history.html', context)

@login_required
def feesstructures(request):
    query = request.GET.get('q')
    fees_structures = FeeStructure.objects.all()

    if query:
        fees_structures = fees_structures.filter(
            Q(class_level__icontains=query) |
            Q(term__icontains=query) |
            Q(amount__icontains=query) |
            Q(effective_date__icontains=query) |
            Q(fee_name__icontains=query)

        ).distinct()

    context = {
        'fees_structures': fees_structures.order_by('-class_level'),
        'query': query,
    }
    return render(request, 'billing/feesstructures.html', context)


@login_required
def invoices(request):
    user = request.user
    query = request.GET.get('q')

    if user.is_staff:
        all_invoices = Invoice.objects.select_related('student', 'student__user').order_by('-due_date')
        students = StudentProfile.objects.all()
        payment_methods = None

    elif user.is_parent:
        all_invoices = Invoice.objects.filter(
            student__parent=user
        ).select_related('student', 'student__user').order_by('-due_date')
        students = StudentProfile.objects.filter(parent=user)
        payment_methods = user.payment_methods.all()

    else:
        profile = StudentProfile.objects.filter(user=user).first()
        if not profile:
            messages.warning(request, "Student profile not found. Please contact admin.")
            return redirect('student_dashboard')

        all_invoices = Invoice.objects.filter(
            student=profile
        ).select_related('student', 'student__user').order_by('-due_date')
        students = [profile]
        payment_methods = None

    if query:
        all_invoices = all_invoices.filter(
            Q(student__user__first_name__icontains=query) |
            Q(student__user__last_name__icontains=query) |
            Q(status__icontains=query) |
            Q(term__icontains=query)
        ).distinct()

    context = {
        'user': user,
        'students': students,
        'methods': payment_methods,
        'all_invoices': all_invoices,
        'query': query,
    }
    return render(request, 'billing/invoices.html', context)

@login_required
def students(request):
    user = request.user
    query = request.GET.get('q')

    if user.is_staff:
        students = StudentProfile.objects.select_related('user').order_by('-class_level')
    else:
        students = StudentProfile.objects.filter(parent=user).select_related('user').order_by('-class_level')

    if query:
        students = students.filter(
            Q(user__first_name__icontains=query) |
            Q(user__last_name__icontains=query) |
            Q(class_level__icontains=query) |
            Q(admission_number__icontains=query)
            
        ).distinct()

    for student in students:
        invoices = student.invoices.all()
        student.paid_count = invoices.filter(status="Fully Paid").count()
        student.unpaid_count = invoices.filter(status__in=["Unpaid", "Partial"]).count()

    context = {
        'students': students,
        'user': user,
        'query': query,
    }
    return render(request, 'billing/students.html', context)

@login_required
def notifications(request):
    notifications = Notification.objects.filter(
        Q(is_public=True) | Q(user=request.user)
    ).distinct().order_by('-created_at')
    
    read_ids = set(UserNotificationStatus.objects.filter(
        user=request.user,
        is_read=True,
        notification__in=notifications
    ).values_list('notification_id', flat=True))
    
    for note in notifications:
        note.is_unread = note.id not in read_ids
    
    unread_count = sum(1 for note in notifications if note.is_unread)
    
    context = {
        'notifications': notifications,
        'unread_count': unread_count,
        'user': request.user,
    }
    return render(request, 'billing/notifications.html', context)


@login_required
def unread_notifications_api(request):
    notifications = Notification.objects.filter(
        Q(is_public=True) | Q(user=request.user)
    ).distinct().order_by('-created_at')

    read_ids = set(UserNotificationStatus.objects.filter(
        user=request.user,
        is_read=True,
        notification__in=notifications
    ).values_list('notification_id', flat=True))

    unread = [
        {
            'id': note.id,
            'title': note.title,
            'message': note.message,
            'created_at': note.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        }
        for note in notifications if note.id not in read_ids
    ]

    return JsonResponse({'unread_notifications': unread})


stripe.api_key = settings.STRIPE_SECRET_KEY

def checkout_page(request):
    return render(request, 'billing/checkout.html')


@csrf_exempt
def create_checkout_session(request, invoice_id):
    if request.method == 'POST':
        invoice = get_object_or_404(Invoice, pk=invoice_id)

        try:
            domain = request.build_absolute_uri('/')[:-1]
            amount_tzs = int(invoice.total_amount * 100)  # Convert Decimal to int (Stripe requires integer)

            if amount_tzs < 50:  # Minimum Stripe amount (adjust as needed)
                messages.error(request, "Amount too low for processing.")
                return redirect('parents_dashboard')  # Update with your actual URL name

            session = stripe.checkout.Session.create(
                payment_method_types=['card'],
                line_items=[{
                    'price_data': {
                        'currency': 'tzs',
                        'product_data': {
                            'name': f'{invoice.student} - {invoice.term}',
                        },
                        'unit_amount': amount_tzs,
                    },
                    'quantity': 1,
                }],
                mode='payment',
                # success_url=f"{domain}/billing/payment_success/?invoice_id={invoice.id}",
                success_url=f"{domain}/billing/payment_status_check/?invoice_id={invoice.id}",
                cancel_url=request.build_absolute_uri('/billing/payment_cancel/'),
                metadata={
                    'user_id': request.user.id,
                    'invoice_id': str(invoice.id),
                    'student_id': str(invoice.student.id),
                    'term': invoice.term,
                    'due_date': str(invoice.due_date),
                }
            )
            return redirect(session.url, code=303)

        except Exception as e:
            messages.error(request, f"Payment error: {str(e)}")
            return redirect('parents_dashboard')

    return JsonResponse({'error': 'Invalid request'}, status=400)


@csrf_exempt
def stripe_webhook(request):
    payload = request.body
    sig_header = request.META.get('HTTP_STRIPE_SIGNATURE')

    try:
        event = stripe.Webhook.construct_event(
            payload, sig_header, settings.STRIPE_WEBHOOK_SECRET
        )
    except (ValueError, stripe.error.SignatureVerificationError):
        return HttpResponse(status=400)

    # Handle successful checkout
    if event['type'] == 'checkout.session.completed':
        session = event['data']['object']
        invoice_id = session['metadata'].get('invoice_id')
        payment_intent_id = session.get('payment_intent')

        if invoice_id and payment_intent_id:
            try:
                invoice = Invoice.objects.get(pk=invoice_id)
                payment_intent = stripe.PaymentIntent.retrieve(payment_intent_id)

                # amount_usd = Decimal(session['amount_total']) / 100  # convert cents to USD
                # amount_tzs = amount_usd * 2770  # convert to TZS
                amount_tzs = Decimal(session['amount_total']) / 100

                payment = Payment.objects.create(
                    invoice=invoice,
                    # amount=Decimal(session['amount_total']) / 100,  # Stripe sends amount in cents

                    amount=amount_tzs.quantize(Decimal('1.'), rounding=ROUND_HALF_UP),
                    payment_date=timezone.now(),
                    method="Card",
                    transaction_id=payment_intent.id
                )

                # Update invoice status based on total payments so far
                total_paid = invoice.payment_set.aggregate(
                    total_paid=Sum('amount')
                )['total_paid'] or Decimal('0')


                # If total_paid is within 5000 TZS of total_amount → Fully Paid; else Partial

                # invoice.total_amount = max(Decimal('0'), invoice.total_amount - total_paid)
                if total_paid >= invoice.total_amount - Decimal('5000'):
                    invoice.status = 'Fully Paid'
                elif total_paid > 0:
                    invoice.status = 'Partial'
                # Else no change if no payment
                invoice.save()

                # Create notification for the user
                Notification.objects.create(
                    user=invoice.student.parent,
                    title="Payment Received",
                    message=(
                        f"We have received a payment of TZS {amount_tzs:,} "
                        f"for invoice #{invoice.id} ({invoice.term}). "
                        f"Current status: {invoice.status}."
                    ),
                    is_public=False
                )
            except Invoice.DoesNotExist:
                pass  # optionally log this
            except Exception as e:
                pass  # optionally log this too

    return HttpResponse(status=200)

@login_required
@parents_required
def payment_success(request):
    invoice_id = request.GET.get('invoice_id')
    if not invoice_id:
        return JsonResponse({'error': 'No invoice ID provided'}, status=400)

    try:
        payment = Payment.objects.filter(invoice_id=invoice_id).latest('payment_date')
    except Payment.DoesNotExist:
        return JsonResponse({'error': 'Payment not found.'}, status=404)
    
    # tzs_amount = payment.amount
    usd_amount = float(payment.amount) / 2770
    context = {
        'amount_usd': usd_amount,
        'amount': payment.amount,
        'transaction_id': payment.transaction_id,
        'payment_method': payment.method,
        'payment_date': payment.payment_date,
    }

    return render(request, 'billing/success.html', context)



@login_required
def payment_status_check(request):
    invoice_id = request.GET.get('invoice_id')
    if not invoice_id:
        return render(request, 'billing/waiting.html', {'error': 'No invoice ID provided'})

    invoice = get_object_or_404(Invoice, pk=invoice_id)

    try:
        payment = Payment.objects.filter(invoice=invoice).latest('payment_date')
        return redirect(f'/billing/payment_success/?invoice_id={invoice_id}')
    except Payment.DoesNotExist:
        return render(request, 'billing/waiting.html', {'invoice': invoice})

def payment_cancel(request):
    return render(request, 'billing/cancel.html')



@login_required
@staff_member_required
def staff_dashboard(request):
    staff = request.user

    recent_payments = Payment.objects.all().order_by('-payment_date')[:5]

    all_pending_invoices = Invoice.objects.filter(
        status__in=['Unpaid', 'Partial']
    ).order_by('-due_date')

    pending_invoices = Invoice.objects.filter(
        status__in=['Unpaid', 'Partial']
    ).order_by('-due_date') 

    all_invoices = Invoice.objects.filter(
        status__in=['Unpaid', 'Partial', 'Fully Paid']
    ).order_by('-due_date') 

    total_due = all_invoices.aggregate(total=Sum('total_amount'))['total'] or 0

    term_paid_total = Payment.objects.filter(
    ).aggregate(total=Sum('amount'))['total'] or 0


    pending_total = total_due - term_paid_total
    if total_due > 0:
        rem_percent = (pending_total / total_due) * 100
    else:
        rem_percent = 0

    pendings_count = Invoice.objects.filter(
        status__in=['Unpaid', 'Partial']
    ).count()

    overdue_count = Invoice.objects.filter(
    status__in=['Unpaid', 'Partial'],
    due_date__lt=now().date()
    ).count()

    unpaid_invoices = Invoice.objects.filter(
        status__in=['Unpaid', 'Partial']
    ).select_related('student')

    upcoming_invoices = unpaid_invoices.order_by('due_date')[:10]  # limit if needed
    all_invoices = Invoice.objects.filter(
    ).order_by('due_date')  # limit if needed
    
    notifications = Notification.objects.filter(Q(is_public=True) | Q(user=request.user)).distinct()
    read_ids = set(UserNotificationStatus.objects.filter(
    user=request.user,
    is_read=True,
    notification__in=notifications
    ).values_list('notification_id', flat=True))

    for note in notifications:
        note.is_unread = note.id not in read_ids

    user_notifications = request.user.notifications.filter(is_public=True)
    has_unread_notifications = any(
    not UserNotificationStatus.objects.filter(user=request.user, notification=note, is_read=True).exists()
    for note in notifications
    )
    unread_notification = UserNotificationStatus.objects.filter(user=request.user, notification__in=notifications, is_read=False).exists()
    unread_count = sum(1 for note in notifications if note.is_unread)
    all_payments = Payment.objects.select_related('invoice', 'invoice__student').order_by('-payment_date')

    context = {
        'user': request.user,
        'students': StudentProfile.objects.order_by('-class_level')[:5],
        'methods': PaymentMethod.objects.all(),
        'recent_payments' : recent_payments,
        'pending_total': pending_total,
        'pendings_count': pendings_count,
        'overdue_count': overdue_count,
        'paid_total': term_paid_total,
        'rem_percent': rem_percent,
        'today': date.today(),
        'upcoming_invoices': upcoming_invoices,
        'all_invoices': all_invoices,
        'all_pending_invoices': all_pending_invoices,
        'notifications': notifications,
        'has_unread_notifications': has_unread_notifications,
        'unread_notification': unread_notification,
        'unread_count': unread_count,
        'fees_structures' : FeeStructure.objects.order_by('-class_level')[:10],
        'all_payments': all_payments,
    }
    return render(request, 'billing/staff_dashboard.html', context)



@login_required
@staff_member_required
def student_detail_dashboard(request, student_id):
    student = get_object_or_404(StudentProfile, id=student_id)
    parent = student.parent

    students = StudentProfile.objects.filter(
        parent = parent
    ).order_by('class_level')

    recent_payments = Payment.objects.filter(
        invoice__student=student
    ).select_related('invoice', 'invoice__student').order_by('-payment_date')[:5]

    all_pending_invoices = Invoice.objects.filter(
        student=student,
        status__in=['Unpaid', 'Partial']
    ).order_by('-due_date')

    all_invoices = Invoice.objects.filter(
        student=student,
        status__in=['Unpaid', 'Partial', 'Fully Paid']
    ).order_by('-due_date')

    pending_invoices = all_pending_invoices[:5]

    # total_due = all_pending_invoices.aggregate(total=Sum('total_amount'))['total'] or 0
    total_due = all_invoices.aggregate(total=Sum('total_amount'))['total'] or 0

    term_paid_total = Payment.objects.filter(
        invoice__student=student,
        # filter by term if needed
    ).aggregate(total=Sum('amount'))['total'] or 0

    pending_total = total_due - term_paid_total
    rem_percent = (pending_total / total_due) * 100 if total_due > 0 else 0

    pendings_count = all_pending_invoices.count()

    overdue_count = Invoice.objects.filter(
        student=student,
        status__in=['Unpaid', 'Partial'],
        due_date__lt=now().date()
    ).count()

    upcoming_invoices = all_pending_invoices.order_by('due_date')[:5]

    notifications = Notification.objects.filter(Q(is_public=True) | Q(user=parent)).distinct()
    read_ids = set(UserNotificationStatus.objects.filter(
        user=parent,
        is_read=True,
        notification__in=notifications
    ).values_list('notification_id', flat=True))

    for note in notifications:
        note.is_unread = note.id not in read_ids

    has_unread_notifications = any(
        not UserNotificationStatus.objects.filter(user=parent, notification=note, is_read=True).exists()
        for note in notifications
    )
    unread_notification = UserNotificationStatus.objects.filter(user=parent, notification__in=notifications, is_read=False).exists()
    unread_count = sum(1 for note in notifications if note.is_unread)

    context = {
        'user': request.user,
        'student': student,
        'parent': parent,
        # 'students': [student],  # Keep template logic consistent
        'students': students,  # Keep template logic consistent
        'methods': parent.payment_methods.all(),
        'recent_payments': recent_payments,
        'pending_total': pending_total,
        'pendings_count': pendings_count,
        'overdue_count': overdue_count,
        'paid_total': term_paid_total,
        'rem_percent': rem_percent,
        'today': date.today(),
        'upcoming_invoices': upcoming_invoices,
        'all_pending_invoices': all_pending_invoices,
        'notifications': notifications,
        'has_unread_notifications': has_unread_notifications,
        'unread_notification': unread_notification,
        'unread_count': unread_count,
        'fee_structure_count': FeeStructure.objects.count(),
        'invoice_count': Invoice.objects.count(),
        'student_count': StudentProfile.objects.count(),
        'payment_count': Payment.objects.count(),
    }

    return render(request, 'billing/student_detail_dashboard.html', context)

@staff_member_required
def create_fee_structure(request):
    form = FeeStructureForm(request.POST or None)
    if form.is_valid():
        form.save()
        messages.success(request, "Fee structure saved successfully")
        return redirect('feesstructures')
    return render(request, 'forms/create_fee_structure.html', {'form': form})

@staff_member_required
def create_invoice(request):
    if request.method == 'POST':
        form = InvoiceForm(request.POST)
        if form.is_valid():
            invoice = form.save(commit=False)

            # Calculate total amount from selected fee structures
            selected_fees = form.cleaned_data['fee_structures']
            total = selected_fees.aggregate(total=Sum('amount'))['total'] or 0
            invoice.total_amount = total
            invoice.save()

            # Manually create InvoiceFee entries (for the through model)
            for fee in selected_fees:
                InvoiceFee.objects.create(
                    invoice=invoice,
                    fee_structure=fee,
                    amount=fee.amount
                )

            messages.success(request, "Invoice created successfully")
            return redirect('invoices') 
    else:
        form = InvoiceForm()

    return render(request, 'forms/create_invoice.html', {'form': form})


@staff_member_required
def create_student_profile(request):
    form = StudentProfileForm(request.POST or None)
    if form.is_valid():
        form.save()
        return redirect('students')
    return render(request, 'forms/create_student_profile.html', {'form': form})

@login_required
@staff_member_required
def export_payments_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="payments.csv"'

    writer = csv.writer(response)
    writer.writerow(['ID', 'Student', 'Amount', 'Method', 'Transaction ID', 'Payment Date'])

    for payment in Payment.objects.select_related('invoice__student'):
        writer.writerow([
            payment.id,
            payment.invoice.student.user.get_full_name(),
            payment.amount,
            payment.method,
            payment.transaction_id,
            payment.payment_date.strftime('%Y-%m-%d %H:%M'),
        ])

    return response

@login_required
@staff_member_required
def export_payments_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"

    ws.append(['ID', 'Student', 'Amount', 'Method', 'Transaction ID', 'Payment Date'])

    for payment in Payment.objects.select_related('invoice__student'):
        ws.append([
            payment.id,
            payment.invoice.student.user.get_full_name(),
            float(payment.amount),
            payment.method,
            payment.transaction_id,
            payment.payment_date.strftime('%Y-%m-%d %H:%M'),
        ])

    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = 'attachment; filename="payments.xlsx"'
    wb.save(response)
    return response

@login_required
def tables(request):
    if request.user.is_staff:
        invoices = Invoice.objects.select_related('student')
    else:
        invoices = Invoice.objects.filter(student__parent=request.user)

    table = InvoiceTable(invoices)
    RequestConfig(request).configure(table)

    return render(request, "forms/tables.html", {"table": table})

@login_required
def student_dashboard(request):
    profile = StudentProfile.objects.filter(user=request.user).first()

    if not profile:
        messages.warning(request, "Please link your parent first.")
        return redirect('link_parent_request')

    if not profile.parent:
        messages.warning(request, "Please link your parent first.")
        return redirect('link_parent_request')
    
    all_pending_invoices = Invoice.objects.filter(
        student=profile,
        status__in=['Unpaid', 'Partial']
    ).order_by('-due_date')

    all_invoices = Invoice.objects.filter(
        student=profile,
        status__in=['Unpaid', 'Partial', 'Fully Paid']
    ).order_by('-due_date')



    pending_invoices = all_pending_invoices[:5]
    # total_due = all_pending_invoices.aggregate(total=Sum('total_amount'))['total'] or 0
    total_due = all_invoices.aggregate(total=Sum('total_amount'))['total'] or 0

    # Payments
    recent_payments = Payment.objects.filter(
        invoice__student=profile
    ).select_related('invoice').order_by('-payment_date')[:5]

    term_paid_total = Payment.objects.filter(
        invoice__student=profile
    ).aggregate(total=Sum('amount'))['total'] or 0

    pending_total = total_due - term_paid_total
    rem_percent = (pending_total / total_due) * 100 if total_due > 0 else 0

    pendings_count = all_pending_invoices.count()

    overdue_count = Invoice.objects.filter(
        student=profile,
        status__in=['Unpaid', 'Partial'],
        due_date__lt=now().date()
    ).count()

    upcoming_invoices = all_pending_invoices.order_by('due_date')[:5]

    # Notifications
    notifications = Notification.objects.filter(Q(is_public=True) | Q(user=request.user)).distinct()
    read_ids = set(UserNotificationStatus.objects.filter(
        user=request.user,
        is_read=True,
        notification__in=notifications
    ).values_list('notification_id', flat=True))

    for note in notifications:
        note.is_unread = note.id not in read_ids

    unread_notification = UserNotificationStatus.objects.filter(
        user=request.user,
        notification__in=notifications,
        is_read=False
    ).exists()

    unread_count = sum(1 for note in notifications if note.is_unread)
    has_unread_notifications = unread_count > 0

    context = {
        'profile': profile,
        'recent_payments': recent_payments,
        'pending_total': pending_total,
        'pendings_count': pendings_count,
        'overdue_count': overdue_count,
        'paid_total': term_paid_total,
        'rem_percent': rem_percent,
        'today': date.today(),
        'upcoming_invoices': upcoming_invoices,
        'all_pending_invoices': all_pending_invoices,
        'notifications': notifications,
        'has_unread_notifications': has_unread_notifications,
        'unread_notification': unread_notification,
        'unread_count': unread_count,
    }
    return render(request, 'billing/student_dashboard.html', context)

@login_required
def link_parent_request(request):
    profile = StudentProfile.objects.filter(user=request.user).first()

    if profile and profile.parent:
        messages.info(request, "You are already linked to a parent.")
        return redirect('student_dashboard')

    if request.method == 'POST':
        form = ParentContactForm(request.POST)
        if form.is_valid():
            email = form.cleaned_data['parent_contact']
            term = form.cleaned_data['term']
            class_level = form.cleaned_data['class_level']
            parent = User.objects.filter(is_parent=True, email=email).first()
            if not parent:
                form.add_error('parent_contact', "No parent account found with that email.")
                messages.error(request, "No parent account found with that email.")
            else:
                # Auto-create profile if not exists
                if not profile:
                    profile = StudentProfile.objects.create(
                        user=request.user,
                        parent_contact=email,
                        admission_number=f"ADM-{request.user.id}",
                        class_level=class_level,
                        term=term,
                    )

                req = ParentRequest.objects.create(
                    student=profile,
                    parent=parent
                )

                note = Notification.objects.create(
                    user=parent,
                    title="Parenting confirmation request",
                    message=(
                        f"{request.user.get_full_name()} wants you to confirm them as your child. Open your email to confirm or reject"
                    ),
                    is_public=False
                )
                UserNotificationStatus.objects.create(user=parent, notification=note, is_read=False)


                accept_url = request.build_absolute_uri(reverse('respond_parent_request', args=[req.id, 'accept']))
                reject_url = request.build_absolute_uri(reverse('respond_parent_request', args=[req.id, 'reject']))

                send_mail(
                    subject="Please confirm your child request",
                    message=f"{request.user.get_full_name()} has listed you as their parent on EduPayy.",
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    recipient_list=[parent.email],
                    html_message=f"""
                <!DOCTYPE html>
                <html>
                <head>
                <meta charset="UTF-8">
                <meta name="viewport" content="width=device-width, initial-scale=1.0">
                </head>
                <body style="font-family: 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; margin: 0; padding: 0; background-color: #f5f7fa; color: #333;">
                <table width="100%" cellpadding="0" cellspacing="0" border="0" style="max-width: 600px; margin: 20px auto; background-color: #ffffff; border-radius: 8px; box-shadow: 0 4px 12px rgba(0, 0, 0, 0.1);">
                    <!-- Header -->
                    <tr>
                    <td style="background: linear-gradient(135deg, #4f46e5, #7c3aed); padding: 30px 20px; border-radius: 8px 8px 0 0; text-align: center;">
                        <h1 style="margin: 0; color: white; font-size: 22px; font-weight: 600;">Parent Confirmation Request</h1>
                    </td>
                    </tr>

                    <!-- Content -->
                    <tr>
                    <td style="padding: 30px 25px;">
                        <p style="font-size: 16px; line-height: 1.6;">Dear {parent.get_full_name()},</p>

                        <p style="font-size: 16px; line-height: 1.6;">
                        <strong>{request.user.get_full_name()}</strong> has listed you as their parent on <strong>EduPayy</strong>.
                        To confirm this relationship, please respond using one of the links below.
                        </p>

                        <div style="text-align: center; margin: 30px 0;">
                        <a href="{accept_url}" style="display: inline-block; background-color: #4f46e5; color: white; text-decoration: none; padding: 12px 24px; border-radius: 6px; font-weight: 500; font-size: 15px; box-shadow: 0 2px 5px rgba(79, 70, 229, 0.3); margin-right: 10px;">Accept Request</a>
                        <a href="{reject_url}" style="display: inline-block; background-color: #e11d48; color: white; text-decoration: none; padding: 12px 24px; border-radius: 6px; font-weight: 500; font-size: 15px; box-shadow: 0 2px 5px rgba(225, 29, 72, 0.3);">Reject Request</a>
                        </div>

                        <p style="font-size: 16px; line-height: 1.6;">
                        If you have any questions or did not expect this request, please ignore it or contact our support team.
                        </p>

                        <p style="font-size: 16px; line-height: 1.6; margin-top: 30px;">Thank you,<br><strong>EduPayy Team</strong></p>
                    </td>
                    </tr>

                    <!-- Footer -->
                    <tr>
                    <td style="padding: 20px 25px; background-color: #f8fafc; border-radius: 0 0 8px 8px; font-size: 14px; color: #64748b; text-align: center;">
                        <p style="margin: 0 0 8px 0;">EduPayy</p>
                        <p style="margin: 0;">Need help? Email us at <a href="mailto:gftinity01@gmail.com" style="color: #4f46e5; text-decoration: none;">gftinity01@gmail.com</a></p>
                    </td>
                    </tr>
                </table>
                </body>
                </html>
                """
                )


                messages.success(request, "Parent request sent! Please wait for their confirmation.")
                return redirect('link_parent_request')
    else:
        form = ParentContactForm()

    return render(request, 'forms/link_parent.html', {'form': form})


@login_required
@parents_required
def respond_parent_request(request, req_id, action):
    req = get_object_or_404(ParentRequest, id=req_id, parent=request.user)

    if req.status != 'pending':
        messages.info(request, f"You already {req.status} this request.")
        return redirect('staff_dashboard' if request.user.is_staff else 'students')

    if action == 'accept':
        # link them
        profile = req.student
        profile.parent = request.user
        profile.save()
        req.status = 'accepted'
        req.save()

        # notify the student
        note = Notification.objects.create(
            user=req.student.user,
            title="Parent confirmed",
            message=f"{request.user.get_full_name()} has accepted you as their child.",
            is_public=False
        )
        UserNotificationStatus.objects.create(user=req.student.user, notification=note, is_read=False)
        messages.success(request, "You’ve confirmed the student relationship.")
    else:
        req.status = 'rejected'
        req.save()
        note = Notification.objects.create(
            user=req.student.user,
            title="Parent request rejected",
            message=f"{request.user.get_full_name()} has rejected your request.",
            is_public=False
        )
        UserNotificationStatus.objects.create(user=req.student.user, notification=note, is_read=False)
        messages.warning(request, "You have rejected the parent request.")

    return redirect('students')




@login_required
@staff_member_required
def create_bulk_invoices(request):
    if request.method == 'POST':
        form = MultiInvoiceForm(request.POST)
        if form.is_valid():
            term = form.cleaned_data['term']
            class_level = form.cleaned_data['class_level']
            due_date = form.cleaned_data['due_date']
            fee_structures = form.cleaned_data['fee_structures']
            all_fitting = form.cleaned_data['all_fitting_students']

            students = (
                StudentProfile.objects.filter(term=term, class_level=class_level)
                if all_fitting else form.cleaned_data['students']
            )

            for student in students:
                invoice = Invoice.objects.create(
                    student=student,
                    term=term,
                    due_date=due_date,
                    total_amount=0  # will update below
                )

                total = 0
                for fee in fee_structures:
                    InvoiceFee.objects.create(
                        invoice=invoice,
                        fee_structure=fee,
                        amount=fee.amount  # explicitly set amount
                    )
                    total += fee.amount

                invoice.total_amount = total
                invoice.save()


            messages.success(request, f"Invoices created for {students.count()} students.")
            return redirect('invoices')
        else:
            return render(request, 'forms/message_error.html', {'form': form})
    else:
        form = MultiInvoiceForm()
    return render(request, 'forms/bulk_create.html', {'form': form})


@login_required
@staff_member_required
def load_invoice_form_options(request):
    term = request.GET.get('term')
    class_level = request.GET.get('class_level')

    students = StudentProfile.objects.filter(term=term, class_level=class_level)
    fee_structures = FeeStructure.objects.filter(term=term, class_level=class_level)

    return render(request, 'forms/partial_form_fields.html', {
        'students': students,
        'fee_structures': fee_structures
    })

